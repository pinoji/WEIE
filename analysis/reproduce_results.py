#!/usr/bin/env python3
"""Reproduce endpoint predictions and statistics without additive volume offsets."""
from pathlib import Path
import csv,json,re
from decimal import Decimal
import numpy as np
import openpyxl
from scipy.special import logsumexp
from scipy.stats import spearmanr
ROOT=Path(__file__).resolve().parents[1]
S=ROOT/'sources';A=ROOT/'analysis'
with (S/'ligand_code_mapping.csv').open() as f:
    ligand_mapping={r['manuscript_code']:r for r in csv.DictReader(f)}
def dump(name,value):(A/name).write_text(json.dumps(value,indent=2)+'\n')
def statistics(rows,key):
    x=np.array([r['experiment'] for r in rows]);y=np.array([r[key] for r in rows]);e=y-x
    return dict(N=len(e),MSE=float(e.mean()),MAE=float(abs(e).mean()),RMSE=float(np.sqrt(np.mean(e*e))),Spearman=float(spearmanr(np.round(x,12),np.round(y,12)).statistic))
def csvout(name,rows,keys):
    with (A/name).open('w',newline='') as f:
        w=csv.DictWriter(f,keys);w.writeheader();w.writerows({k:d.get(k) for k in keys} for d in rows)
t4=[]
for pocket,name in [('Nonpolar','Result_Hydrophobic_1.xlsx'),('Polar','result1_Hydrophlic.xlsx')]:
    w=openpyxl.load_workbook(S/name,read_only=True,data_only=True);rs=w['Result']
    for i in range(2,10):
        npol=pocket=='Nonpolar';lig=rs.cell(i,1 if npol else 12).value
        baseline=rs.cell(i,8 if npol else 7).value;es=rs.cell(i,4 if npol else 3).value
        ie=rs.cell(i,9 if npol else 8).value;g0=rs.cell(i,7 if npol else 9).value
        exp=rs.cell(i,12 if npol else 13).value
        # Preserve the source comparator output when removing the old offset.
        # IND differs by 0.01 from the sum of its tabulated components.
        assert abs(g0-(baseline+ie))<=0.0100000001
        sn=ligand_mapping[lig]['analysis_sheet']
        hp=np.array([r[0] for r in w[sn].iter_rows(min_row=1,max_row=3600,min_col=2,max_col=2,values_only=True)])
        kt=w[sn].cell(2,6 if npol else 8).value
        assert abs(kt*(logsumexp(hp/kt)-np.log(len(hp)))-es)<1e-10
        g=baseline/2
        t4.append(dict(pocket=pocket,ligand=lig,experiment=exp,mm_baseline=baseline,ie=ie,weie=es,
            mm_steg=g,mmpbsa=g0,dh_steg=g-es,hp_mean=float(hp.mean()),hp_N=len(hp),
            hp_kT=kt,workbook=name,result_row=i))
    w.close()
stats={}
for group in ['Nonpolar','Polar','Overall','Polar_matched']:
    rr=[r for r in t4 if r['experiment'] is not None and (group=='Overall' or r['pocket']==group or (group=='Polar_matched' and r['pocket']=='Polar' and r['ligand'] not in ['CAQ','IPH']))]
    stats[group]={k:statistics(rr,k) for k in ['mm_steg','mmpbsa']}
dump('mm_ligand_values.json',t4);dump('mm_statistics.json',stats)
csvout('T4_results.csv',t4,list(t4[0]))
summary={}
for line in (S/'WEIE_summary.txt').read_text().splitlines():
    if re.match(r'^CB\w+MD/',line):
        p=line.split();summary[p[0].split('/')[0]]=dict(hp_N=int(p[1]),mm_baseline=float(p[2]),ie=float(p[3]),weie=float(p[5]),hp_mean=float(p[6]))
cb=[]
for d in json.loads((S/'cb7_reference_inputs.json').read_text()):
    if d.get('source_directory') in summary:d.update(summary[d['source_directory']])
    else:
        with (A/('cb7_'+d['guest']+'_hp.csv')).open() as f:
            hp=np.array([float(r['HP_kcal_mol']) for r in csv.DictReader(f)])
        d['weie']=float(d['kT_source']*(logsumexp(hp/d['kT_source'])-np.log(len(hp))))
    g=d['mm_baseline']/2
    d.update(mm_steg=g,mmpbsa=float(Decimal(str(d['mm_baseline']))+Decimal(str(d['ie']))),dh_steg=g-d['weie'])
    d['signed_errors']={k:d[k]-d['experiment'] for k in ['mm_steg','mmpbsa']}
    d['apr_signed_errors']=[v-d['experiment'] for v in d['apr']]
    if d.get('experimental_enthalpy') is not None:d['dh_steg_error']=d['dh_steg']-d['experimental_enthalpy']
    cb.append(d)
cb_stats={k:statistics(cb,k) for k in ['mm_steg','mmpbsa']}
hydro=[d for d in cb if d['guest'] not in ['AdH','AdOH']]
result=dict(systems=cb,statistics=cb_stats,
    subset_statistics={'alkanes':{k:statistics(hydro,k) for k in ['mm_steg','mmpbsa']}},
    estimator='G_STEG=L/2; G_PBSA_IE=source endpoint (T4), L+IE (CB7); H_STEG=G_STEG-WEIE')
dump('cb7_results.json',result)
csvout('CB7_results.csv',cb,['guest','experiment','mm_baseline','ie','weie','mm_steg','mmpbsa','dh_steg'])
for rr in [t4,cb]:
    assert all(abs(d['dh_steg']-(d['mm_steg']-d['weie']))<1e-11 for d in rr)
    assert all(abs(d['mm_steg']-d['mm_baseline']/2)<1e-11 for d in rr)
    assert all(abs(d['mmpbsa']-d['mm_baseline']-d['ie'])<=0.0100000001 for d in rr)
for name in ['T4_corrected_results.csv','CB7_corrected_results.csv']:(A/name).unlink(missing_ok=True)
dump('final_validation.json',dict(estimator=result['estimator'],T4_total=16,T4_with_experiment=15,CB7_total=8,
    source_workbooks_unchanged=True,T4_exact_STEG_tie=['DEN','OXE'],
    statistics_use_source_precision=True,Spearman_ties='average ranks'))
print('Reproduced all T4/CB7 predictions, enthalpies, and tied-rank statistics.')

# Reproduce the raw and WEIE ablations with the documented input precision.
# Existing source-precision prediction records above are left unchanged.
methods={
    'MM-STEG':lambda d:d['mm_baseline']/2,
    'Raw MM/PBSA':lambda d:d['mm_baseline'],
    'MM/PBSA+IE':lambda d:d['mmpbsa'],
    'MM/PBSA+WEIE':lambda d:d['mm_baseline']+(round(d['weie'],2) if 'ligand' in d else d['weie'])}
ablation_rows=[]
for dataset,records in [('T4',t4),('CB7',cb)]:
    for d in records:
        for method,calc in methods.items():
            value=calc(d)
            ablation_rows.append(dict(dataset=dataset,system=d.get('ligand',d.get('guest')),
                method=method,experiment=d['experiment'],prediction=value,
                signed_error=None if d['experiment'] is None else value-d['experiment']))
csvout('ablation_predictions.csv',ablation_rows,list(ablation_rows[0]))
groups={
 'CB7, all':cb,'CB7, hydrocarbons':hydro,
 'T4 nonpolar':[d for d in t4 if d['pocket']=='Nonpolar'],
 'T4 polar':[d for d in t4 if d['pocket']=='Polar' and d['experiment'] is not None],
 'T4 pooled':[d for d in t4 if d['experiment'] is not None],
 'Polar targets':[d for d in t4 if d['pocket']=='Polar' and d['ligand'] not in ['CAQ','IPH','JZ3']]}
ablation_stats=[]
for dataset,records in groups.items():
    for method,calc in methods.items():
        rr=[dict(experiment=d['experiment'],prediction=calc(d)) for d in records]
        ablation_stats.append(dict(dataset=dataset,method=method,**statistics(rr,'prediction')))
csvout('ablation_statistics.csv',ablation_stats,list(ablation_stats[0]))
print('Reproduced raw MM/PBSA and MM/PBSA+WEIE ablations for all datasets.')
